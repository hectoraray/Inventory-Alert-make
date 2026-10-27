"""Cross-check the spreadsheet formulas against the reference JavaScript implementation.

The alert logic exists twice in this portfolio:
  * inventory-alerts-n8n/n8n/01_classify_stock.js   (reference, with its own unit tests)
  * the helper-column formulas built by scripts/build_template.py (this project)

This script feeds ~1,000 input combinations (numbers as text, blanks, negatives, padded
spaces, every previous status...) through BOTH and fails if any result differs.

Requirements: LibreOffice (`soffice`) and Node.js on the PATH, and the sibling project
`inventory-alerts-n8n` next to this one.

Usage:
    python scripts/verify_against_n8n.py
"""
from __future__ import annotations

import itertools
import json
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent))
from build_template import build_workbook  # noqa: E402

ROOT = Path(__file__).resolve().parents[1]
REFERENCE_JS = ROOT.parent / "inventory-alerts-n8n" / "n8n" / "01_classify_stock.js"

NODE_RUNNER = """
const fs = require("fs");
const src = fs.readFileSync(process.argv[2], "utf8");
const rows = JSON.parse(fs.readFileSync(0, "utf8"));
const $input = { all: () => rows.map((json) => ({ json })) };
const out = new Function("$input", src)($input).map((i) => i.json);
console.log(JSON.stringify(out.map((r) => ({
  status: r.status, previous_status: r.previous_status,
  should_notify: r.should_notify ? 1 : 0, needs_update: r.needs_update ? 1 : 0,
}))));
"""

STOCKS = ["120", "34", "21", "20", "11", "10", "1", "0", 0, 5, 12, "  5 ", "", "   ", "n/a", "abc", "-3", None]
REORDERS = ["20", "40", 25, "0", "-5", "x", ""]
LASTS = ["", "OK", "LOW", "CRITICAL", "OUT", "DATA_ERROR", "  low ", "foo"]


def main() -> int:
    for tool in ("soffice", "node"):
        if shutil.which(tool) is None:
            print(f"ERROR: '{tool}' not found on PATH")
            return 2
    if not REFERENCE_JS.exists():
        print(f"ERROR: reference implementation not found at {REFERENCE_JS}")
        return 2

    rows, js_rows = [], []
    for i, (stock, reorder, last) in enumerate(itertools.product(STOCKS, REORDERS, LASTS)):
        rows.append([f"S{i}", "Product", "Category", stock, reorder, "Supplier", last])
        js_rows.append({"sku": f"S{i}", "product_name": "Product", "category": "Category",
                        "stock": stock, "reorder_point": reorder, "supplier": "Supplier",
                        "last_alert_status": last})

    with tempfile.TemporaryDirectory() as tmp:
        tmp_dir = Path(tmp)
        source = tmp_dir / "combos.xlsx"
        build_workbook(rows, source, with_help_sheet=False)

        # LibreOffice recalculates every formula while converting
        out_dir = tmp_dir / "recalculated"
        subprocess.run(["soffice", "--headless", "--convert-to", "xlsx", "--outdir", str(out_dir), str(source)],
                       check=True, capture_output=True, timeout=300)
        ws = load_workbook(out_dir / "combos.xlsx", data_only=True)["inventory"]

        runner = tmp_dir / "run_js.js"
        runner.write_text(NODE_RUNNER)
        js = json.loads(subprocess.run(["node", str(runner), str(REFERENCE_JS)], input=json.dumps(js_rows),
                                       capture_output=True, text=True, check=True).stdout)

    col = {ws.cell(1, c).value: c for c in range(1, ws.max_column + 1)}
    fields = ["status", "previous_status", "should_notify", "needs_update"]
    mismatches = []
    for i, expected in enumerate(js):
        got = {f: ws.cell(i + 2, col[f]).value for f in fields}
        if got != expected:
            mismatches.append((rows[i][3:7], got, expected))

    print(f"Compared {len(rows)} input combinations")
    if mismatches:
        print(f"FAILED: {len(mismatches)} mismatches. First ones:")
        for inputs, got, expected in mismatches[:5]:
            print(f"  inputs={inputs}\n    sheet={got}\n    js   ={expected}")
        return 1
    print("OK: spreadsheet formulas and the JavaScript reference agree on every case")
    return 0


if __name__ == "__main__":
    sys.exit(main())
